"""Student Management System: a Tkinter desktop app that stores student marks in Excel."""

import os
from tkinter import *
from tkinter import messagebox, ttk

from openpyxl import Workbook, load_workbook

FILE_NAME = "students.xlsx"
HEADERS = ["Roll", "Name", "Sub1", "Sub2", "Sub3", "Total", "Average", "Grade"]


def get_grade(avg):
    if avg >= 90:
        return "A+"
    if avg >= 75:
        return "A"
    if avg >= 60:
        return "B"
    if avg >= 50:
        return "C"
    if avg >= 35:
        return "D"
    return "F"


def load_records():
    if not os.path.exists(FILE_NAME):
        return []
    sheet = load_workbook(FILE_NAME).active
    return [list(row) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None]


def save_records(records):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    sheet.append(HEADERS)
    for record in records:
        sheet.append(record)
    workbook.save(FILE_NAME)


root = Tk()
root.title("Student Management System")
root.geometry("820x560")
root.configure(bg="#f4f6fb")

Label(root, text="STUDENT MARKS DETAILS", font=("Arial", 18, "bold"), bg="#f4f6fb", fg="#2c3e50").pack(pady=10)

form = Frame(root, bg="#f4f6fb")
form.pack()

labels = ["Roll No", "Student Name", "Subject 1", "Subject 2", "Subject 3", "Total", "Average", "Grade"]
entries = []
for i, text in enumerate(labels):
    Label(form, text=text, bg="#f4f6fb", font=("Arial", 10)).grid(row=i // 2, column=(i % 2) * 2, sticky=W, padx=10, pady=5)
    e = Entry(form, width=22)
    e.grid(row=i // 2, column=(i % 2) * 2 + 1, padx=10, pady=5)
    entries.append(e)

e_roll, e_name, e_s1, e_s2, e_s3, e_total, e_avg, e_grade = entries
for e in (e_total, e_avg, e_grade):
    e.configure(state="readonly")


def set_readonly(entry, value):
    entry.configure(state=NORMAL)
    entry.delete(0, END)
    entry.insert(0, value)
    entry.configure(state="readonly")


def refresh_table():
    table.delete(*table.get_children())
    for record in load_records():
        table.insert("", END, values=record)


def clear_form():
    for e in entries:
        set_readonly(e, "") if e in (e_total, e_avg, e_grade) else e.delete(0, END)
    e_roll.focus()


def calculate():
    roll = e_roll.get().strip()
    name = e_name.get().strip()
    if not roll or not name:
        messagebox.showwarning("Missing data", "Please enter Roll No and Student Name.")
        return

    try:
        marks = [int(e.get()) for e in (e_s1, e_s2, e_s3)]
    except ValueError:
        messagebox.showerror("Invalid marks", "Marks must be whole numbers.")
        return
    if any(m < 0 or m > 100 for m in marks):
        messagebox.showerror("Invalid marks", "Marks must be between 0 and 100.")
        return

    total = sum(marks)
    avg = round(total / 3, 2)
    grade = get_grade(avg)

    set_readonly(e_total, total)
    set_readonly(e_avg, avg)
    set_readonly(e_grade, grade)

    # Update the student if the roll number already exists, otherwise add a new row
    records = [r for r in load_records() if str(r[0]) != roll]
    records.append([roll, name, *marks, total, avg, grade])
    save_records(records)
    refresh_table()

    messagebox.showinfo("Saved", f"Data for {name} saved to {FILE_NAME}")


def delete_selected():
    selected = table.selection()
    if not selected:
        messagebox.showwarning("No selection", "Select a student in the table first.")
        return
    roll = str(table.item(selected[0])["values"][0])
    if messagebox.askyesno("Delete", f"Delete record for Roll No {roll}?"):
        save_records([r for r in load_records() if str(r[0]) != roll])
        refresh_table()


def on_select(_event):
    selected = table.selection()
    if not selected:
        return
    values = table.item(selected[0])["values"]
    clear_form()
    for e, v in zip((e_roll, e_name, e_s1, e_s2, e_s3), values[:5]):
        e.insert(0, v)
    set_readonly(e_total, values[5])
    set_readonly(e_avg, values[6])
    set_readonly(e_grade, values[7])


buttons = Frame(root, bg="#f4f6fb")
buttons.pack(pady=10)
Button(buttons, text="CALCULATE & SAVE", command=calculate, width=18, bg="#27ae60", fg="white").grid(row=0, column=0, padx=6)
Button(buttons, text="CLEAR", command=clear_form, width=12, bg="#7f8c8d", fg="white").grid(row=0, column=1, padx=6)
Button(buttons, text="DELETE SELECTED", command=delete_selected, width=16, bg="#c0392b", fg="white").grid(row=0, column=2, padx=6)

table_frame = Frame(root)
table_frame.pack(fill=BOTH, expand=True, padx=15, pady=10)
table = ttk.Treeview(table_frame, columns=HEADERS, show="headings", height=10)
for h in HEADERS:
    table.heading(h, text=h)
    table.column(h, width=95, anchor=CENTER)
scroll = ttk.Scrollbar(table_frame, orient=VERTICAL, command=table.yview)
table.configure(yscrollcommand=scroll.set)
table.pack(side=LEFT, fill=BOTH, expand=True)
scroll.pack(side=RIGHT, fill=Y)
table.bind("<<TreeviewSelect>>", on_select)

refresh_table()
root.mainloop()
